"""Build canonical dataset from INE Nomenclátor XLSX."""

from __future__ import annotations

import json
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

from .admin import PROVINCE_NAMES, PROVINCE_TO_COMMUNITY, community_code_for_province
from .ine_codes import (
    ParsedIneCode,
    UnitType,
    build_ine_code,
    entity_code_for,
    is_eligible_singular_entity,
    parse_ine_code,
)
from .normalize import normalize_for_search


DATASET_SOURCE = "INE_NOMENCLATOR"
DEFAULT_REFERENCE_DATE = "2025-01-01"


@dataclass
class RawRow:
    ine_code: str
    name: str
    population_total: int
    population_male: int
    population_female: int
    parsed: ParsedIneCode


def _parse_unit_field(unit_field: str) -> tuple[str, str]:
    """XLSX 'Unidad Poblacional' column: 'CCSSNN NAME'."""
    unit_field = unit_field.strip()
    if len(unit_field) < 7:
        raise ValueError(f"Unidad poblacional too short: {unit_field!r}")
    suffix = unit_field[:6]
    name = unit_field[7:].strip() if len(unit_field) > 7 else ""
    if not suffix.isdigit():
        raise ValueError(f"Invalid unit suffix: {unit_field!r}")
    return suffix, name


def read_ine_xlsx(path: Path) -> Iterator[RawRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    if header[:3] != ("Provincia", "Municipio", "Unidad Poblacional"):
        raise ValueError(f"Unexpected XLSX header: {header}")

    for row in rows:
        if not row or row[0] is None:
            continue
        province = str(row[0]).zfill(2)
        municipality = str(row[1]).zfill(3)
        unit_suffix, name = _parse_unit_field(str(row[2]))
        cc, ss, nn = unit_suffix[0:2], unit_suffix[2:4], unit_suffix[4:6]
        ine_code = build_ine_code(province, municipality, cc, ss, nn)
        parsed = parse_ine_code(ine_code)
        yield RawRow(
            ine_code=ine_code,
            name=name,
            population_total=int(row[3]),
            population_male=int(row[4]),
            population_female=int(row[5]),
            parsed=parsed,
        )
    workbook.close()


def build_canonical(
    xlsx_path: Path,
    output_dir: Path,
    reference_date: str = DEFAULT_REFERENCE_DATE,
) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    municipalities: dict[str, dict[str, Any]] = {}
    places: list[dict[str, Any]] = []
    nuclei: list[dict[str, Any]] = []
    seen_place_ids: set[str] = set()

    for raw in read_ine_xlsx(xlsx_path):
        parsed = raw.parsed
        province_code = parsed.province_code
        municipality_id = parsed.municipality_ine_code

        if parsed.unit_type == UnitType.MUNICIPALITY_TOTAL:
            municipalities[municipality_id] = {
                "id": municipality_id,
                "name": raw.name,
                "normalizedName": normalize_for_search(raw.name),
                "provinceCode": province_code,
                "provinceName": PROVINCE_NAMES[province_code],
                "autonomousCommunityCode": community_code_for_province(province_code),
                "autonomousCommunityName": PROVINCE_TO_COMMUNITY[province_code],
                "population": raw.population_total,
            }
            continue

        if parsed.unit_type == UnitType.SINGULAR_ENTITY:
            place_id = raw.ine_code
            if place_id in seen_place_ids:
                raise ValueError(f"Duplicate entidad singular id: {place_id}")
            seen_place_ids.add(place_id)
            eligible = is_eligible_singular_entity(raw.population_total, parsed.unit_type)
            places.append(
                {
                    "id": place_id,
                    "ineCode": place_id,
                    "name": raw.name,
                    "normalizedName": normalize_for_search(raw.name),
                    "placeType": UnitType.SINGULAR_ENTITY.value,
                    "population": raw.population_total,
                    "populationMale": raw.population_male,
                    "populationFemale": raw.population_female,
                    "populationReferenceDate": reference_date,
                    "municipalityId": municipality_id,
                    "provinceCode": province_code,
                    "provinceName": PROVINCE_NAMES[province_code],
                    "autonomousCommunityCode": community_code_for_province(province_code),
                    "autonomousCommunityName": PROVINCE_TO_COMMUNITY[province_code],
                    "collectiveCode": parsed.collective_code,
                    "singularCode": parsed.singular_code,
                    "latitude": None,
                    "longitude": None,
                    "coordinateSource": None,
                    "source": DATASET_SOURCE,
                    "sourceVersion": reference_date,
                    "isEligible": eligible,
                }
            )
            continue

        if parsed.unit_type == UnitType.NUCLEUS:
            parent = entity_code_for(parsed)
            nuclei.append(
                {
                    "id": raw.ine_code,
                    "ineCode": raw.ine_code,
                    "parentEntityCode": parent,
                    "name": raw.name,
                    "normalizedName": normalize_for_search(raw.name),
                    "placeType": UnitType.NUCLEUS.value,
                    "population": raw.population_total,
                    "populationReferenceDate": reference_date,
                    "municipalityId": municipality_id,
                    "latitude": None,
                    "longitude": None,
                }
            )

    _write_jsonl(output_dir / "canonical_places.jsonl", places)
    _write_jsonl(output_dir / "municipalities.jsonl", list(municipalities.values()))
    _write_jsonl(output_dir / "nuclei.jsonl", nuclei)

    eligible_places = [p for p in places if p["isEligible"]]
    summary = {
        "datasetVersion": reference_date,
        "populationReferenceDate": reference_date,
        "source": DATASET_SOURCE,
        "importedAt": datetime.now(timezone.utc).isoformat(),
        "sourceFile": str(xlsx_path.name),
        "counts": {
            "singularEntities": len(places),
            "eligiblePlaces": len(eligible_places),
            "ineligibleSingularOverThreshold": sum(
                1 for p in places if not p["isEligible"] and p["population"] >= 10_000
            ),
            "municipalities": len(municipalities),
            "nuclei": len(nuclei),
        },
    }
    (output_dir / "summary.json").write_text(
        json.dumps(summary, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return summary


def _write_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8") as handle:
        for record in records:
            handle.write(json.dumps(record, ensure_ascii=False) + "\n")
